from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.models.requests import InferenceRequest, InferenceResponse


class ArtifactStore:
    """Writes human-readable request inputs and generated outputs separately."""

    def __init__(self, root: str) -> None:
        self.root = Path(root)
        self.inputs = self.root / "inputs"
        self.outputs = self.root / "outputs"
        self.workloads = self.root / "workloads"
        self.results = self.root / "results"
        self.inputs.mkdir(parents=True, exist_ok=True)
        self.outputs.mkdir(parents=True, exist_ok=True)
        self.workloads.mkdir(parents=True, exist_ok=True)
        self.results.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _safe(value: str | None) -> str:
        return "".join(char if char.isalnum() or char in "-_" else "_" for char in (value or "standalone"))

    def write_input(self, request_id: str, payload: InferenceRequest) -> str:
        experiment = self._safe(payload.experiment_id)
        folder = self.inputs / experiment
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{request_id}.txt"
        content = payload.prompt or "\n".join(f"[{message.role}] {message.content}" for message in payload.messages or [])
        path.write_text(f"request_id: {request_id}\nmodel: {payload.model}\npriority: {payload.priority}\nmax_output_tokens: {payload.max_output_tokens}\n\n{content}\n", encoding="utf-8")
        return str(path)

    def write_output(self, request_id: str, experiment_id: str | None, result: InferenceResponse) -> str:
        experiment = self._safe(experiment_id)
        folder = self.outputs / experiment
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{request_id}.txt"
        path.write_text(f"request_id: {request_id}\nmodel: {result.model}\ninput_tokens: {result.input_tokens}\noutput_tokens: {result.output_tokens}\nttft_ms: {result.ttft_ms}\n\n{result.output}\n", encoding="utf-8")
        return str(path)

    def write_workload_excel(self, pair_id: str, requests: list[InferenceRequest]) -> str:
        path = self.workloads / f"{self._safe(pair_id)}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "LLM Inputs"
        sheet.append(["row", "size", "input", "predicted_output_tokens", "model", "input_source", "prediction_source"])
        for index, payload in enumerate(requests, 1):
            sheet.append([index, payload.metadata.get("size", ""), payload.prompt or "", payload.expected_output_tokens, payload.model, payload.metadata.get("input_source", ""), payload.metadata.get("output_prediction_source", "")])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["C"].width = 100
        workbook.save(path)
        return str(path)

    def append_result_excel(self, request_id: str, experiment_id: str | None, payload: InferenceRequest, result: InferenceResponse, backend: str | None, latency_ms: float) -> str:
        path = self.results / f"{self._safe(experiment_id)}.xlsx"
        if path.exists():
            workbook = load_workbook(path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "LLM Results"
            sheet.append(["request_id", "size", "policy", "model", "backend", "input", "output", "input_tokens", "predicted_output_tokens", "actual_output_tokens", "ttft_ms", "latency_ms"])
            sheet.freeze_panes = "A2"
        sheet.append([request_id, payload.metadata.get("size", ""), payload.scheduler_policy or "", result.model, backend or "", payload.prompt or "", result.output, result.input_tokens, payload.expected_output_tokens, result.output_tokens, result.ttft_ms, latency_ms])
        sheet.column_dimensions["F"].width = 80
        sheet.column_dimensions["G"].width = 100
        workbook.save(path)
        return str(path)
